"""Informe de revisión de la orden, con estados explícitos y original anexo."""
import json,hashlib
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]


def run():
    'IA: Crea copia de revisión con evidencia real, digests completos y orden original; nunca marca B/C como cerradas ni modifica el original.'
    data=json.loads((ROOT/'reports/iteration_contracts.json').read_text(encoding='utf8'))
    expected=json.loads((ROOT/'reports/expected.json').read_text(encoding='utf8'))['cases']
    actual=json.loads((ROOT/'reports/modular.json').read_text(encoding='utf8'))
    doc=Document();sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
    style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(5)
    for name in ('Title','Heading 1','Heading 2'):doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    footer=sec.footer.paragraphs[0];footer.text='Ruinas · iteration-A1-2026-10-24 · revisión del equipo · '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_heading('Revisión de la orden de trabajo de Ruinas',0)
    doc.add_paragraph('Iteración A1 · 24 de octubre de 2026 · Commit 2bcd7be · develop')
    doc.add_paragraph('Resultado: contratos A aplicados parcialmente y reserva común de escalera integrada. La orden completa NO está terminada. La fase B sigue abierta y la fase C no está aceptada. Este documento permite revisar lo realizado y continuar sin confundir infraestructura con validación estructural.')
    doc.add_heading('Estado de los requisitos',1)
    table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for c,t in zip(table.rows[0].cells,['Punto','Estado','Resultado o límite']):c.text=t
    rows=[('A1','Aplicado con C','IDs del catálogo, density_ref, roles y familias; nuevas densidades provisionales. Multiplicador artístico debe aplicarse sobre catálogo cuando exista el puente.'),('A2','Parcial','model_scale declarado; adaptador convierte medidas. Falta generación paramétrica independiente a dos escalas y prueba de tensiones de apoyo.'),('A3','Aplicado','RECHAZADO / PROVISIONAL / ACEPTADO. El camino probado sigue PROVISIONAL.'),('A4','Aplicado en esquema','Lista de contactos por pareja, medidas None y validación negativa. No se han medido contactos reales del edificio.'),('A5','Aplicado','Eje de vano, fachada y planta en ID; carpintería separada. El piso es atributo, no anfitrión del hueco.'),('A6','Actualizado','Mapa, changelog e histórico v2. El recuento actual es 35 pruebas puras.'),('B','Parcial','Reserva previa compartida y caché verificadas. Faltan elementos de vigueta/cabecero/tablero gobernados por plan y aceptación de pérdida de apoyo.'),('C','No aceptado','Auditoría del trabajador realizada; no se ha cambiado el umbral ni declarado equivalencia con Bullet.')]
    for row in rows:
        for c,t in zip(table.add_row().cells,row):c.text=t
    doc.add_paragraph('No se añadieron nuevas familias de fachada, daño ni física de edificio completo. No se elevó ningún dato a confianza B. No se sustituyó ningún ZIP anterior.')
    doc.add_page_break();doc.add_heading('Pruebas y evidencia',0)
    doc.add_paragraph('35 pruebas unitarias puras aprobadas. Incluyen referencias y ciclos inválidos, capacidades desconocidas, aceptación con y sin fuente, área negativa, reparto de carga mayor que uno, fusión de contactos e identidad estable al insertar vanos. No son 35 simulaciones Bullet.')
    doc.add_paragraph('Los cuatro casos de geometría conservaron sus digests. La integración en tres plantas y la restauración de construction_reservations desde caché pasaron. La captura de esta entrega es de inspección general; no prueba todavía un encuentro estructural aceptado.')
    doc.add_heading('Escala y variabilidad',1)
    r=data['scale_mesh_test'][0];v=data['variability']
    doc.add_paragraph('La prueba mide una viga generada, representada a 1:30 y 1:60. Ambas mediciones dan %.6f kg y %.3f mm² de superficie. Comprueba la frontera de conversión, pero no el requisito completo A2: no genera dos edificios con parámetros independientes ni calcula tensión de asiento. Peso/superficie total no es tensión estructural de apoyo.'%(r['mass_kg'],r['area_mm2']))
    doc.add_paragraph('Variabilidad: 1.000 semillas, CV solicitado 0,25, media %.6f, desviación %.6f, mínimo %.6f y máximo %.6f. Distribución lognormal truncada en su variable normal; la truncación modifica la media y dispersión respecto a la distribución original.'%(v['mean'],v['stdev'],v['min'],v['max']))
    doc.add_heading('Valores C y procedencia',1)
    doc.add_paragraph('Camino crítico de prueba TIMBER_SOFT: densidad, compresión y apoyo perpendicular; 3 campos C, ninguno B. Las densidades añadidas son revoco 1.200, vidrio 2.500 y carpintería ligera 600 kg/m³: valores provisionales de trabajo, sin ensayo ni norma. La carpintería mixta requiere descomposición posterior. Vidrio no tiene aún physical_role ni familia de resistencia soportados por el laboratorio; no se asigna por aproximación a piedra.')
    doc.add_paragraph('Catálogo materials-0.2.0-C. Firma SHA256: '+data['materials_signature'])
    doc.add_heading('Auditoría física',1)
    doc.add_paragraph('El trabajador existente todavía usa PHYSICS_MASS_SCALE=10.000 y PHYSICS_SOLVER_SCALE=100. Escala gravedad y umbrales dentro de Bullet; el umbral original deriva de la masa del par. Esto NO satisface la fase C basada en material y área. No se sustituye por model_scale sin calibración. No se pasaron capacity_N ni resistencias MPa directamente a breaking_threshold.')
    doc.add_page_break();doc.add_heading('Inspección y siguientes criterios de cierre',0)
    doc.add_picture(str(ROOT/'reports/iteration_reservations.png'),width=Cm(11.5))
    doc.add_paragraph('Vista general de la casa generada en la prueba. No es una captura del fallo ni evidencia de comparación planificador–Bullet.',style='Caption')
    doc.add_paragraph('Para cerrar A2: producir dos representaciones paramétricas independientes del mismo edificio y medir fuerzas y áreas de asiento. Para cerrar B: declarar viguetas, cabeceros, tablero y extremos de apoyo antes de las mallas; hacer que los consumidores lean ese plan; impedir reparaciones silenciosas en edificios dañados. El bastidor alternativo heredado de floors sigue pendiente de ese tratamiento.')
    doc.add_paragraph('Después: probar pérdida de apoyo redundante, pérdida de ambos apoyos y conservación de elementos ajenos. Solo con B aceptada iniciar C: uniones por catálogo/contacto, varias semillas y BALANCED/PRECISE, conjunto de caídas igual al planificador y captura de encuentro y resultado.')
    doc.add_paragraph('Revisión del archivo: estructura DOCX, tablas e imagen comprobadas por lectura. La paginación visual en Word/PDF no se ha podido verificar automáticamente: el entorno no dispone de Word ni LibreOffice. Se conserva la orden original en el anexo, sin modificar el archivo recibido.')
    doc.add_page_break();doc.add_heading('Digests geométricos y archivos de evidencia',0)
    for row in actual:
        key=row['case'];before=expected[key]['digest'];after=row['digest'];assert before==after
        doc.add_heading(key,2)
        doc.add_paragraph('Antes: '+before+'\nDespués: '+after+'\nResultado: conservado.')
    doc.add_paragraph('Evidencia en el repositorio: reports/iteration_A1_check.log; iteration_A_validation.log; iteration_B_reservations.log; iteration_reservation_integration.log; iteration_contracts.json; iteration_contracts.log. Pruebas: tests/test_construction_contracts.py, test_construction_core.py y blender_iteration_contracts.py.')
    doc.add_page_break();doc.add_heading('Anexo con la orden original',0)
    original=Document(r'C:\Users\yhora\Downloads\Orden de trabajo Ruinas, siguiente iteración.docx')
    for p in original.paragraphs:
        if p.text:doc.add_paragraph(p.text)
    out=ROOT/'dist/Ruinas_Orden_revision_iteration_A1_2026-10-24.docx';doc.save(out)
    check=Document(out);assert len(check.tables)==1 and len(check.inline_shapes)==1
    manifest={'delivery':'iteration-A1-2026-10-24','commit':'2bcd7be','catalog':'materials-0.2.0-C','docx_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'source_docx_sha256':hashlib.sha256(Path(r'C:\Users\yhora\Downloads\Orden de trabajo Ruinas, siguiente iteración.docx').read_bytes()).hexdigest(),'complete_order':False,'phase_B_accepted':False,'phase_C_accepted':False,'rendered_docx':False,'geometry_digests':{r['case']:r['digest'] for r in actual}}
    (ROOT/'reports/iteration_A1_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
    print(out,flush=True)


if __name__=='__main__':run()
